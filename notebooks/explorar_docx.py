# Importamos las librerías
# Módulo de rutas de Python
from pathlib import Path
# Importamos solo Document que nos representa el documento entero
from docx import Document
# Counter para obtener el diccionario de frecuencia de estilos de párrafo en el texto
from collections import Counter

# Vamos hasta la raíz
RAIZ = Path(__file__).resolve().parent.parent
# Path donde se encuentra la tesis en formato docx
RUTA = RAIZ / "data" / "raw" / "electronic_and_optical_properties_of_organic_molecules_at_metal_surfaces_studied_by_scanning_tunneling_microscopy_oscar_jover_arrate_phd_thesis.docx"

# Comprobación de si existe el archivo
print(f"Buscando en: {RUTA}")
print(f"¿Existe? {RUTA.exists()}\n")

# Abrimos el docx para descomprimir el archivo y obtenerlo en formato XML
doc = Document(RUTA)

# Listamos todos los parrafos, es decir el número de parrafos que hay
print(f"Total de párrafos: {len(doc.paragraphs)}")
# Lo mismo, número de tablas que tiene la tesis
print(f"Total de tablas: {len(doc.tables)}\n")

# Contamos los estilos que hay de texto y cuantas veces aparecen los printeamos por frecuencia de aparición de más a menos común
estilos = Counter(p.style.name for p in doc.paragraphs)
print("Estilos encontrados:")
for estilo, cuenta in estilos.most_common():
    print(f"  {estilo}: {cuenta}")