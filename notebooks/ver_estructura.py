
# Módulo de rutas de Python
from pathlib import Path
# Importamos Document que nos representa el documento entero
from docx import Document

# Vamos hasta la raíz del archivo. Este script solo funciona si está en una carpeta un nivel por debajo de la raíz
RAIZ = Path(__file__).resolve().parent.parent
# Path donde está el archivo .docx de Word con la tesis
RUTA = RAIZ / "data" / "raw" / "electronic_and_optical_properties_of_organic_molecules_at_metal_surfaces_studied_by_scanning_tunneling_microscopy_oscar_jover_arrate_phd_thesis.docx"

# Leemos el word. Básicamente Document() lee el documento y lo carga en doc como un objeto abierto en XML
doc = Document(RUTA)

# Representamos los niveles de título según el siguiente diccionario
NIVELES = {"Heading 1": 0, "Heading 2": 1, "Heading 3": 2}

# Con vacios calculamos el número de líneas vacías puestas en la tesis para separar entre párrafos
vacios = 0

# Vamos párrafo a párrafo
for p in doc.paragraphs:
    # Quitamos espacios y saltos de línea al principio y al final. Si nos da True, es decir, tenemos una línea vacía
    # sumamos uno al número de líneas vacías (var vacios)
    if not p.text.strip():
        vacios += 1
    # Saca el nivel, coge el tipo de Heading que es la key y nos da el value de nuestro diccionario. Luego nos separa los
    # headings por sangrías para mejor visualización. Si no hay headings directamente nos da None y lo descartamos con el continue
    nivel = NIVELES.get(p.style.name)
    if nivel is None:
        continue
    print("    " * nivel + p.text.strip())

# Indicamos también el número de párrafos vacíos
print(f"\n--- Párrafos vacíos: {vacios} ---")